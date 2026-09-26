"""Build the editable workbook from the authoritative Markdown lab files.

Requires python-docx. Render the DOCX with the documents skill's render_docx.py
and inspect every page before replacing the companion PDF.
"""
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path(__file__).resolve().parents[1]
LABS = ROOT / "VLA_Drone_Starter_Kit/labs"
OUTPUT = ROOT / "Guide/VLA_Drone_Lab_Workbook.docx"
LINKS = {path.name: f"lab{path.name[:2]}" for path in LABS.glob("[0-9][0-9]_*.md")}
LINKS["WORKSHEET.md"] = "worksheet"
INLINE = re.compile(r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)")


def add_link(paragraph, label, target):
    link = OxmlElement("w:hyperlink")
    if target in LINKS:
        link.set(qn("w:anchor"), LINKS[target])
    else:
        link.set(qn("r:id"), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    colour = OxmlElement("w:color"); colour.set(qn("w:val"), "003399")
    props.append(colour); run.append(props)
    text = OxmlElement("w:t"); text.text = label
    run.append(text); link.append(run); paragraph._p.append(link)


def inline(paragraph, text):
    for part in INLINE.split(text):
        if part.startswith("["):
            match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
            if match:
                add_link(paragraph, *match.groups())
                continue
        run = paragraph.add_run(part[2:-2] if part.startswith("**") else
                                part[1:-1] if part.startswith("`") else part)
        if part.startswith("**"): run.bold = True
        if part.startswith("`"):
            run.font.name = "Courier New"; run.font.size = Pt(9)


def bookmark(paragraph, name, identifier):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(identifier)); start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), str(identifier))
    paragraph._p.insert(0, start); paragraph._p.append(end)


def table(document, rows):
    rows = [row for row in rows if not re.fullmatch(r"[\s|:\-]+", row)]
    values = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in rows]
    result = document.add_table(rows=len(values), cols=len(values[0]))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    count = len(values[0])
    proportions = {2: [.36, .64], 3: [.23, .37, .40], 4: [.08, .28, .27, .37]}[count]
    for column, fraction in zip(result.columns, proportions): column.width = Inches(6.6 * fraction)
    properties = result._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{side}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn(f"w:{key}"), value)
        borders.append(border)
    properties.append(borders)
    for i, row in enumerate(values):
        trpr = result.rows[i]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if i == 0: trpr.append(OxmlElement("w:tblHeader"))
        for j, value in enumerate(row):
            cell = result.cell(i, j)
            cell.width = Inches(6.6 * proportions[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "bottom", "left", "right"):
                element = OxmlElement(f"w:{side}")
                element.set(qn("w:w"), "95"); element.set(qn("w:type"), "dxa")
                margins.append(element)
            tcpr.append(margins)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E7EDF5" if i == 0 else "FFFFFF")
            tcpr.append(shading)
            paragraph = cell.paragraphs[0]
            paragraph.style = "Table text"
            inline(paragraph, value)
            if i == 0:
                for run in paragraph.runs: run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def append_markdown(document, path, index):
    lines = path.read_text().splitlines()
    if path.name == "README.md":
        lines = lines[:lines.index("## Reference sources")]
    elif path.name == "WORKSHEET.md":
        reference = (LABS / "README.md").read_text().split("## Reference sources", 1)[1]
        lines += ["", "## Reference sources", *reference.splitlines()]
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip(): i += 1; continue
        if line.startswith("```"):
            code = []; i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i]); i += 1
            paragraph = document.add_paragraph("\n".join(code), "Code")
            paragraph.paragraph_format.keep_together = True
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i]); i += 1
            table(document, rows); continue
        elif line.startswith("# "):
            paragraph = document.add_paragraph(line[2:], "Title" if index == 0 else "Heading 1")
            if index: paragraph.paragraph_format.page_break_before = True
            if path.name in LINKS: bookmark(paragraph, LINKS[path.name], index + 1)
            if index == 0:
                document.add_paragraph("Mac and workstation edition   |   26 September 2026", "Subtitle")
        elif line.startswith("## "):
            paragraph = document.add_paragraph(line[3:], "Heading 2")
            breaks = {
                "08_commands.md": "Milestone B Implement the bridge",
                "10_smolvla_training.md": "Inspect the training bundle",
                "11_mission_manager.md": "Test faults before flight integration",
                "12_unified_vla.md": "Milestone C Integrate proposals with the manager",
            }
            if breaks.get(path.name) == line[3:]:
                paragraph.paragraph_format.page_break_before = True
        elif re.match(r"^(- |\d+\. )", line):
            match = re.match(r"^(- |\d+\. )(.*)", line)
            paragraph = document.add_paragraph(style="List Bullet" if match[1] == "- " else "Normal")
            inline(paragraph, match[2] if match[1] == "- " else line)
        else:
            inline(document.add_paragraph(), line)
        i += 1


def main():
    document = Document()
    for border in document.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    section = document.sections[0]
    section.page_width = Inches(8.27); section.page_height = Inches(11.69)
    section.top_margin = Inches(.6); section.bottom_margin = Inches(.6)
    section.left_margin = Inches(.83); section.right_margin = Inches(.83)
    section.header_distance = Inches(.28); section.footer_distance = Inches(.3)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"; normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.0
    for name, size in (("Title", 25), ("Heading 1", 19), ("Heading 2", 12), ("Subtitle", 11)):
        style = document.styles[name]
        style.font.name = "Arial"; style.font.size = Pt(size); style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(7 if name == "Heading 2" else 0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True
    code = document.styles.add_style("Code", 1)
    code.font.name = "Courier New"; code.font.size = Pt(8.2)
    code.paragraph_format.line_spacing = 1.05
    code.paragraph_format.space_after = Pt(8)
    table_style = document.styles.add_style("Table text", 1)
    table_style.base_style = normal
    table_style.font.size = Pt(9)
    table_style.paragraph_format.space_after = Pt(0)
    table_style.paragraph_format.line_spacing = 1.05
    header = section.header.paragraphs[0]
    header.text = "VLA DRONE LAB WORKBOOK"
    header.runs[0].font.name = "Arial"; header.runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Lab workbook   |   ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    paths = [LABS / "README.md", *sorted(LABS.glob("[0-9][0-9]_*.md")), LABS / "WORKSHEET.md"]
    for index, path in enumerate(paths): append_markdown(document, path, index)
    document.core_properties.title = "VLA drone lab workbook"
    document.core_properties.subject = "Practical instructions and completion evidence for twelve labs"
    document.core_properties.author = "VLA Drone Student Package"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
